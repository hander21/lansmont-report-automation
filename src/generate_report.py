"""
Generates a draft Word report from the template using table-cell token replacement.

Text tokens   {{ field_name }}        are replaced with parsed values.
Image tokens  [IMG: slot_name]        are replaced with photos.
Chart tokens  [CHART: slot_name]      are replaced with chart images.
"""

import logging
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches

from src.insert_images import get_fit_dimensions
from src.utils import sanitize_for_filename

logger = logging.getLogger(__name__)

_TEXT_TOKEN = re.compile(r"\{\{\s*(\w+)\s*\}\}")
_IMAGE_TOKEN = re.compile(r"\[(IMG|CHART):\s*(\S+?)\s*\]")


def build_template_context(parsed: dict, cfg: dict) -> dict:
    """Return a copy of parsed values plus the draft watermark."""
    ctx = dict(parsed)
    ctx["draft_watermark"] = "** DRAFT — NOT FOR DISTRIBUTION **"
    return ctx


def _build_image_map(discovered: dict, output_paths: dict, cfg: dict) -> dict:
    """
    Map slot names to resolved image Paths.

    Photo slots:  pre_1..pre_7, accel_1..accel_2, seq2_1..seq2_8,
                  seq3_1..seq3_4, seq4_1..seq4_4, seq6_1..seq6_4,
                  seq7_1..seq7_4, post_1..post_10
    Chart slots:  derived from filename via _chart_slot_name()
    """
    image_map: dict[str, Path] = {}

    # ── Photo folders → slot prefix ──
    photo_dirs = [
        ("photos_pre", "pre"),
        ("photos_accel", "accel"),
        ("photos_seq2", "seq2"),
        ("photos_seq3", "seq3"),
        ("photos_seq4", "seq4"),
        ("photos_seq6", "seq6"),
        ("photos_seq7", "seq7"),
        ("photos_post", "post"),
    ]
    for key, prefix in photo_dirs:
        d = output_paths.get(key)
        if d and Path(d).exists():
            for i, p in enumerate(sorted(Path(d).glob("*")), start=1):
                image_map[f"{prefix}_{i}"] = p

    # ── Charts ──
    charts_dir = output_paths.get("charts")
    if charts_dir and Path(charts_dir).exists():
        for chart_path in sorted(Path(charts_dir).glob("*")):
            slot = _chart_slot_name(chart_path.name)
            if slot:
                image_map[slot] = chart_path
            else:
                logger.warning("Chart not matched to any slot: %s", chart_path.name)

    return image_map


def _chart_slot_name(filename: str) -> str | None:
    """Derive a slot key from a chart filename, e.g. CHART_SEQ3_ROT_DROP_1_A.png → seq3_rot_drop_1_a"""
    name = filename.upper()
    if not name.startswith("CHART_"):
        return None
    stem = name[len("CHART_"):].replace(".PNG", "").replace(".JPG", "").replace(".JPEG", "")
    return stem.lower()


def _replace_text_in_cell(cell, ctx: dict) -> list[str]:
    """Replace all {{ field }} tokens in a cell. Returns list of replaced field names."""
    replaced = []
    for para in cell.paragraphs:
        full_text = "".join(r.text for r in para.runs)
        if "{{" not in full_text:
            continue
        new_text = full_text
        for match in _TEXT_TOKEN.finditer(full_text):
            field = match.group(1)
            value = ctx.get(field, "")
            new_text = new_text.replace(match.group(0), str(value))
            replaced.append(field)
        if new_text != full_text:
            for i, run in enumerate(para.runs):
                run.text = new_text if i == 0 else ""
    return replaced


def _replace_image_in_cell(cell, image_map: dict, chart_cfg: dict, photo_cfg: dict) -> dict | None:
    """
    If a cell contains [IMG: slot] or [CHART: slot], replace with the image.
    Returns an info dict if an image was inserted, else None.
    """
    for para in cell.paragraphs:
        full_text = "".join(r.text for r in para.runs)
        if "[IMG:" not in full_text and "[CHART:" not in full_text:
            continue
        m = _IMAGE_TOKEN.search(full_text)
        if not m:
            continue
        kind, slot = m.group(1), m.group(2).lower()
        image_path = image_map.get(slot)

        if not image_path or not image_path.exists():
            logger.warning("Image slot [%s: %s] has no matching file.", kind, slot)
            return None

        is_chart = kind == "CHART"
        max_w = chart_cfg["max_width_inches"] if is_chart else photo_cfg["max_width_inches"]
        max_h = chart_cfg["max_height_inches"] if is_chart else photo_cfg["max_height_inches"]
        w, _ = get_fit_dimensions(image_path, max_w, max_h)

        # Clear all runs, then insert picture into the first run
        for run in para.runs:
            run.text = ""
        para.runs[0].add_picture(str(image_path), width=Inches(w))

        return {"kind": kind, "slot": slot, "file": image_path.name}

    return None


def generate_draft_report(
    parsed: dict,
    discovered: dict,
    output_paths: dict,
    cfg: dict,
) -> tuple[Path, list]:
    """
    Open the .docx template, replace all tokens in table cells with values and images,
    and save the draft report. Returns (report_path, inserted_images).
    """
    template_path = Path(cfg["template_path"])
    doc = Document(str(template_path))

    ctx = build_template_context(parsed, cfg)
    image_map = _build_image_map(discovered, output_paths, cfg)

    chart_cfg = cfg["chart_settings"]
    photo_cfg = cfg["photo_settings"]

    inserted_images: list[dict] = []
    replaced_text_fields: set[str] = set()

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                # Try image replacement first
                result = _replace_image_in_cell(cell, image_map, chart_cfg, photo_cfg)
                if result:
                    inserted_images.append(result)
                    continue
                # Then text replacement
                fields = _replace_text_in_cell(cell, ctx)
                replaced_text_fields.update(fields)

    # Build output filename
    customer = sanitize_for_filename(parsed["customer_name"])
    project = sanitize_for_filename(parsed["project_number"])
    date = sanitize_for_filename(parsed["test_date"])
    test_type = sanitize_for_filename(parsed["test_type"])
    filename = f"DRAFT_{customer}_{project}_{test_type}_{date}.docx"
    output_file = output_paths["final_report"] / filename

    doc.save(str(output_file))
    logger.info("Draft report saved: %s", output_file)
    logger.info(
        "Replaced %d text field(s), inserted %d image(s).",
        len(replaced_text_fields),
        len(inserted_images),
    )

    return output_file, inserted_images
